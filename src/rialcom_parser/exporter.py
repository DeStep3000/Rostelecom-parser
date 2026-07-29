from collections.abc import Sequence
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from rialcom_parser.models import Tariff

_HEADERS: Final = (
    "Название тарифа",
    "Количество каналов",
    "Скорость доступа, Мбит/с",
    "Абонентская плата, руб.",
)


class ExcelTariffExporter:
    """Сохраняет тарифы в читаемый Excel-документ."""

    def export(self, tariffs: Sequence[Tariff], output_path: Path) -> None:
        """Создает XLSX-файл и родительский каталог при необходимости."""
        output_path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Тарифы"

        worksheet.append(_HEADERS)
        for tariff in tariffs:
            worksheet.append(
                (
                    tariff.name,
                    tariff.channels,
                    tariff.speed_mbps,
                    tariff.monthly_fee,
                )
            )

        self._format_worksheet(worksheet)
        workbook.save(output_path)

    @staticmethod
    def _format_worksheet(worksheet: Worksheet) -> None:
        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.column_dimensions["A"].width = 62
        worksheet.column_dimensions["B"].width = 22
        worksheet.column_dimensions["C"].width = 25
        worksheet.column_dimensions["D"].width = 27

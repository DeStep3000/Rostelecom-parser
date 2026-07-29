from pathlib import Path

from openpyxl import load_workbook

from rialcom_parser.exporter import ExcelTariffExporter
from rialcom_parser.models import Tariff


def test_export_tariffs_to_excel(tmp_path: Path) -> None:
    output_path = tmp_path / "nested" / "tariffs.xlsx"
    tariffs = [
        Tariff("РиалКомий-15+", None, 200, 1000),
        Tariff("Комбо Лайт + РиалКом Интернет 50 + ТВ", 165, 50, 529),
    ]

    ExcelTariffExporter().export(tariffs, output_path)

    workbook = load_workbook(output_path)
    worksheet = workbook["Тарифы"]

    assert worksheet.max_row == 3
    assert worksheet["A1"].value == "Название тарифа"
    assert worksheet["B2"].value is None
    assert worksheet["C2"].value == 200
    assert worksheet["D3"].value == 529
    assert worksheet.freeze_panes == "A2"
